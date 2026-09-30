# scripts/taco-preparo/gerar.py
# Extrai do PDF da TACO (4ª edição revisada e ampliada, NEPA/Unicamp, 2011) o Apêndice 2 -- "Protocolos padronizados para os alimentos
# preparados" (Dra. Sonia Tucunduva Philippi, FSP/USP): COMO cada alimento preparado foi feito para a análise (quanto óleo, tempo de
# cozimento, ingredientes). Os valores da TACO valem para ESSE preparo.
#   -> prisma/data/taco-preparo.json   e   (opcional) src/data/tacoPreparo.js do app
# Uso (na pasta do fitos-final):  python3 scripts/taco-preparo/gerar.py caminho/taco_4_edicao.pdf [caminho/do/app/src/data/tacoPreparo.js]
# Requer: pip install pymupdf openpyxl
import json, os, re, sys, unicodedata
import pymupdf as fitz
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.join(HERE, '..', '..')
PDF = sys.argv[1]; OUT_JS = sys.argv[2] if len(sys.argv) > 2 else None
OUT_JSON = os.path.join(ROOT, 'prisma', 'data', 'taco-preparo.json')

def norm(s):
    s = unicodedata.normalize('NFD', s.lower()); s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

# nomes oficiais da TACO (planilha, já com os nomes corrigidos de scripts/taco-correcoes)
corr = json.load(open(os.path.join(ROOT, 'scripts', 'taco-correcoes', 'correcoes.json'), encoding='utf-8'))
REN = {c['de']: c['para'] for c in corr['nomes']}
ws = load_workbook(os.path.join(ROOT, 'prisma', 'Taco-4a-Edicao.xlsx'), data_only=True)['CMVCol taco3']
NAMES = {}
for r in ws.iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], (int, float)) and isinstance(r[1], str):
        nm = re.sub(r'\s+', ' ', r[1]).strip(); nm = REN.get(nm, nm); NAMES[norm(nm)] = nm
# títulos do livro que não são exatamente o nome do alimento na tabela
ALIASES = {'Pipoca': 'Pipoca, com óleo de soja, sem sal'}

d = fitz.open(PDF)
pages = [d[i].get_text() for i in range(len(d))]
start = next(i for i, t in enumerate(pages) if 'Apêndice 2. Protocolos' in t and i > 60)
end = next(i for i in range(start + 1, len(pages)) if re.search(r'(?m)^\s*Índice Remissivo', pages[i], re.I))
lines = []
for i in range(start, end):
    for l in pages[i].split('\n'):
        s = l.strip()
        if re.fullmatch(r'\d{1,3}', s) or s.startswith('Apêndice 2') or s.startswith('Dra. Sonia'): continue
        lines.append(s)
while lines and lines[0] in ('', 'Abreviações', 'g', 'gramas', 'mL mililitro', 'L', 'litro', 'kg quilograma'): lines.pop(0)

PROC = [i for i, l in enumerate(lines) if re.fullmatch(r'Procedimentos?:', l)]
def find_title(pi):
    j = pi - 1
    while j >= 0 and lines[j] == '': j -= 1
    cand = lines[j]
    if norm(cand) in NAMES or cand in ALIASES: return j
    for k in range(j, max(-1, j - 50), -1):            # receitas regionais: a lista de ingredientes vem entre o título e "Procedimentos:"
        if lines[k] and (norm(lines[k]) in NAMES or lines[k] in ALIASES): return k
    return j
TITLES = [find_title(pi) for pi in PROC]

def parse_body(body):
    steps, sections, paras, obs = [], [], [], []
    cur_list = steps; cur = None
    for l in body:
        if l == '': cur = None if not cur else cur; continue
        if re.match(r'^Observação:', l): obs.append(l[len('Observação:'):].strip()); cur = ('obs', obs); continue
        m = re.match(r'^(\d+)\.\s*(.*)$', l)
        if m: cur = {'n': m.group(1), 't': m.group(2)}; cur_list.append(cur); continue
        if re.fullmatch(r'[A-ZÁÉÍÓÚÂÊÔÃÕÇ0-9][^.:]{1,45}:', l):
            sec = {'title': l[:-1], 'steps': []}; sections.append(sec); cur_list = sec['steps']; cur = None; continue
        if isinstance(cur, tuple): cur[1][-1] += ' ' + l; continue
        if isinstance(cur, dict): cur['t'] = (cur['t'] + ' ' + l).strip(); continue
        if l.startswith('- '): cur = {'n': '', 't': l[2:]}; cur_list.append(cur); continue
        cur = {'n': '', 't': l}; cur_list.append(cur)        # texto corrido (receitas regionais sem numeração)
    clean = lambda L: [{'n': s['n'], 't': re.sub(r'\s+', ' ', s['t']).replace('È necessário', 'É necessário').strip()} for s in L if s['t'].strip()]   # (erro de digitação do livro)
    return clean(steps), [{'title': s['title'], 'steps': clean(s['steps'])} for s in sections], ' '.join(obs).strip()

# protocolos em fases ("1ª fase: Cocção" / "2ª fase: Fritura") têm vários "Procedimentos:" sob o mesmo título
def phase_line(pi):
    j = pi - 1
    while j >= 0 and lines[j] == '': j -= 1
    return j if re.match(r'^\d+ª fase', lines[j]) else None
PHASE = [phase_line(pi) for pi in PROC]
groups = []
for k, ti in enumerate(TITLES):
    if groups and groups[-1][0] == ti: groups[-1][1].append(k)
    else: groups.append((ti, [k]))

out, unmatched = [], []
for gi, (ti, ks) in enumerate(groups):
    nxt = groups[gi + 1][0] if gi + 1 < len(groups) else len(lines)
    raw_title = lines[ti]; name = ALIASES.get(raw_title) or NAMES.get(norm(raw_title))
    if not name: unmatched.append(raw_title); continue
    e = {'name': name}
    if len(ks) == 1:
        pi = PROC[ks[0]]
        ingredients = [re.sub(r'\s+', ' ', l) for l in lines[ti + 1:pi] if l]
        steps, sections, obs = parse_body(lines[pi + 1:nxt])
        if ingredients: e['ingredients'] = ingredients
        if steps: e['steps'] = steps
        if sections: e['sections'] = sections
        if obs: e['obs'] = obs
    else:
        secs, obs_all = [], []
        for idx, k in enumerate(ks):
            end = PHASE[ks[idx + 1]] if idx + 1 < len(ks) else nxt
            steps, sections, obs = parse_body(lines[PROC[k] + 1:end])
            secs.append({'title': lines[PHASE[k]] if PHASE[k] is not None else f'Fase {idx + 1}', 'steps': steps})
            for sec in sections: secs.append({'title': sec['title'], 'steps': sec['steps']})
            if obs: obs_all.append(obs)
        e['sections'] = secs
        if obs_all: e['obs'] = ' '.join(obs_all)
    out.append(e)
names = [e['name'] for e in out]
dups = sorted({n for n in names if names.count(n) > 1})
print(f'{len(PROC)} protocolos no livro -> {len(out)} ligados a alimentos da TACO; sem par: {unmatched}; repetidos: {dups}')
empty = [e['name'] for e in out if not e.get('steps') and not e.get('sections')]
print('sem passos:', empty)
os.makedirs(os.path.dirname(OUT_JSON), exist_ok=True)
json.dump({'fonte': 'TACO 4ª ed. rev. e ampl. (NEPA/Unicamp, 2011), Apêndice 2 - Protocolos padronizados para os alimentos preparados (Dra. Sonia Tucunduva Philippi, FSP/USP)', 'protocolos': out},
          open(OUT_JSON, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('->', OUT_JSON, os.path.getsize(OUT_JSON) // 1024, 'KB')
if OUT_JS:
    body = json.dumps({norm(e['name']): {k: v for k, v in e.items()} for e in out}, ensure_ascii=False, separators=(',', ':'))
    with open(OUT_JS, 'w', encoding='utf-8') as f:
        f.write('// src/data/tacoPreparo.js — GERADO por fitos-final/scripts/taco-preparo/gerar.py (não edite à mão).\n')
        f.write('// Como a TACO (4ª ed., Apêndice 2) preparou cada alimento preparado para a análise. Chave = nome do alimento sem acento/pontuação, minúsculo.\n')
        f.write('export const TACO_PREPARO_FONTE = "TACO 4ª ed. rev. e ampl. (NEPA/Unicamp, 2011), Apêndice 2 - Protocolos padronizados para os alimentos preparados (Dra. Sonia Tucunduva Philippi, FSP/USP)";\n')
        f.write('export const TACO_PREPARO = ' + body + ';\n')
    print('->', OUT_JS, os.path.getsize(OUT_JS) // 1024, 'KB')
